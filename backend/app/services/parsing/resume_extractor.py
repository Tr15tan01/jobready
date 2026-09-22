"""
Extracts plain text from an uploaded resume file.

Per section 9/29: the original file is treated as untrusted input, is
never persisted after extraction, is validated for size/type before any
processing, and never touches a privileged AI system prompt directly —
callers must pass the extracted text through the resume-extraction
*prompt template*, not embed it into arbitrary instructions.
"""
import io

from docx import Document
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB
ALLOWED_CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
}


async def read_and_validate_upload(file: UploadFile) -> tuple[bytes, str]:
    if file.content_type not in ALLOWED_CONTENT_TYPES:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Only PDF, DOCX, or TXT resumes are supported.",
        )
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File exceeds the 10MB limit.",
        )
    if len(data) == 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Empty file.")
    return data, ALLOWED_CONTENT_TYPES[file.content_type]


def extract_text(data: bytes, kind: str) -> str:
    if kind == "pdf":
        return _extract_pdf(data)
    if kind == "docx":
        return _extract_docx(data)
    if kind == "txt":
        return _extract_txt(data)
    raise ValueError(f"Unsupported file kind: {kind}")


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    if not text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Could not extract text from this PDF (it may be a scanned image).",
        )
    return text


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    text = "\n".join(paragraphs).strip()
    if not text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Could not extract text from this document.",
        )
    return text


def _extract_txt(data: bytes) -> str:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="ignore")
    text = text.strip()
    if not text:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Empty text file.")
    return text
